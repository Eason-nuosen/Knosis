"""SpreadsheetBench scoring: port of the cell comparison in upstream `evaluation/evaluation.py`; formulas are recalculated with LibreOffice before scoring.

Deliverable `predictions.jsonl`, one `{"id": instruction id, "answer": {"<case no>": "<output xlsx path>"}}` per line,
paths relative to the deliverable's directory. Per case: the output file is opened, recalculated and re-saved by LibreOffice (as in upstream `open_spreadsheet.py`,
so openpyxl `data_only=True` sees formula results), then values are compared cell by cell within `answer_position`:
numbers rounded to 2 decimals, datetime to integer Excel serial, time to the minute, numeric strings compared as numbers, empty string equals empty, different types are unequal.
Two differences from upstream: answer files containing TODAY() / NOW() are recalculated too (cached values are from the annotation date, otherwise correct formulas would fail);
if the `answer_position` sheet name has no exact match, match ignoring spaces and case (upstream data has `Sheet 2` vs `Sheet2`).
Metrics: soft = mean per-instruction case pass rate, hard = share of instructions passing all cases. Upstream's color comparison is commented out; likewise we compare neither color nor format.
"""

from __future__ import annotations

import datetime
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from . import answers, mean, register

SOFFICE = ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice")


def find_soffice() -> str | None:
    for candidate in [os.environ.get("KNOSIS_SOFFICE", ""), *SOFFICE]:
        if candidate and (shutil.which(candidate) or Path(candidate).is_file()):
            return shutil.which(candidate) or candidate
    return None


def recalc(files, outdir) -> dict[Path, Path]:
    """Open with LibreOffice and save as xlsx (recalculating formulas, writing cached values); returns {original: recalculated file}, failures omitted.
    Each call uses its own user profile directory, so multiple processes can run concurrently."""
    soffice, outdir = find_soffice(), Path(outdir)
    if soffice is None:
        raise RuntimeError("LibreOffice not found (set KNOSIS_SOFFICE); SpreadsheetBench scoring needs it to recalculate formulas")
    out = {}
    with tempfile.TemporaryDirectory() as profile:
        for i, src in enumerate(files):
            target = outdir / str(i)
            target.mkdir(parents=True, exist_ok=True)
            subprocess.run([soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--calc",
                            "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(target), str(src)],
                           capture_output=True, timeout=180)
            made = target / (Path(src).stem + ".xlsx")
            if made.exists():
                out[Path(src)] = made
    return out


def excel_serial(dt) -> float:
    delta = dt - datetime.datetime(1899, 12, 30)
    return delta.days + delta.seconds / 86400.0


def transform(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    if isinstance(v, datetime.datetime):
        return round(excel_serial(v), 0)
    if isinstance(v, datetime.time):
        return str(v)[:-3]
    if isinstance(v, str):
        try:
            return round(float(v), 2)
        except ValueError:
            return v
    return v


def same(a, b) -> bool:
    a, b = transform(a), transform(b)
    if a in ("", None) and b in ("", None):
        return True
    return type(a) is type(b) and a == b


def col_num(name) -> int:
    n = 0
    for c in name:
        n = n * 26 + ord(c) - 64
    return n


def col_name(n) -> str:
    out = ""
    while n:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out


def cells(rng, max_row=None) -> list[str]:
    """Range -> cell names. Handles several notations in upstream data: lowercase columns, full-width colons, whole columns (`A:G`, rows up to `max_row`),
    end without a column (`BD2:308` means the same column)."""
    rng = rng.upper().replace("：", ":")
    if ":" not in rng:
        return [rng]
    parse = lambda x: ("".join(ch for ch in x if ch.isalpha()), "".join(ch for ch in x if ch.isdigit()))  # noqa: E731
    (c1, r1), (c2, r2) = parse(rng.split(":")[0]), parse(rng.split(":")[1])
    c2, r1, r2 = c2 or c1, int(r1 or 1), int(r2 or max_row or 1)
    return [f"{col_name(c)}{r}" for c in range(col_num(c1), col_num(c2) + 1) for r in range(r1, r2 + 1)]


def regions(position, default_sheet) -> list[tuple[str, str]]:
    """`answer_position` -> [(sheet, range)]; in upstream data sheet-name quotes are sometimes unbalanced."""
    out = []
    for part in position.split(","):
        if "!" in part:
            sheet, rng = part.split("!", 1)
            sheet = sheet.strip().strip("'")
        else:  # default to the first sheet, name kept verbatim (some sheet names have leading/trailing spaces)
            sheet, rng = default_sheet, part
        out.append((sheet, rng.strip().strip("'").replace("$", "")))
    return out


def sheet_name(name, names):
    """Exact sheet name match; some upstream answer_position names differ from the file (e.g. `Sheet 2` vs `Sheet2`),
    so without an exact match, take the unique match ignoring spaces and case."""
    if name in names:
        return name
    key = lambda x: x.replace(" ", "").lower()  # noqa: E731
    hits = [n for n in names if key(n) == key(name)]
    return hits[0] if len(hits) == 1 else None


VOLATILE = ("TODAY(", "NOW(")


def volatile(path) -> bool:
    """Answer file contains TODAY() / NOW(): cached values are from the annotation date, so it must be recalculated on the same day as the output to be comparable."""
    import openpyxl

    wb = openpyxl.load_workbook(path)
    return any(isinstance(c.value, str) and any(f in c.value.upper() for f in VOLATILE)
               for ws in wb.worksheets for row in ws.iter_rows() for c in row)


def compare(gold_file, proc_file, position) -> tuple[bool, str, str]:
    """(match, detail without the expected value, detail with the expected value)."""
    import openpyxl

    try:
        gold = openpyxl.load_workbook(gold_file, data_only=True)
        proc = openpyxl.load_workbook(proc_file, data_only=True)
    except Exception as error:  # noqa: BLE001 - failing to open counts as wrong
        return False, f"cannot open output: {type(error).__name__}", f"cannot open output: {error}"
    for sheet, rng in regions(position, gold.sheetnames[0]):
        g, p = sheet_name(sheet, gold.sheetnames), sheet_name(sheet, proc.sheetnames)
        if p is None:
            return False, f"worksheet {sheet!r} not found", f"worksheet {sheet!r} not found"
        if g is None:
            return False, f"answer worksheet {sheet!r} missing upstream", ""
        wg, wp = gold[g], proc[p]
        for name in cells(rng, wg.max_row):
            if not same(wg[name].value, wp[name].value):
                return (False, f"{sheet}!{name} differs (output {wp[name].value!r})",
                        f"{sheet}!{name}: expected {wg[name].value!r}, output {wp[name].value!r}")
    return True, "", ""


def score_cases(task, outputs: dict, scratch) -> list[dict]:
    """`outputs`：{case no: output file path or None}."""
    given = {n: Path(p) for n, p in outputs.items() if p and Path(p).is_file()}
    golds = [g for g in task.answer if volatile(g)] if given else []
    fixed = recalc(list(given.values()) + golds, scratch) if given else {}
    rows = []
    for n, gold in zip(task.meta["cases"], task.answer):
        src = given.get(n)
        if src is None:
            rows.append({"case": n, "ok": False, "detail": "missing output file", "gold": ""})
            continue
        ok, detail, full = compare(fixed.get(Path(gold), gold), fixed.get(src, src), task.meta["answer_position"])
        if src not in fixed:
            detail += " (LibreOffice recalculation failed, compared the original file)"
        rows.append({"case": n, "ok": ok, "detail": detail, "gold": full})
    return rows


@register("spreadsheetbench")
def score(ws, state, ds, path, ids) -> dict:
    tasks = {t.id: t for t in ds.tasks}
    preds = answers(path) if path and Path(path).exists() else {}
    items, reference = [], []
    with tempfile.TemporaryDirectory() as scratch:
        for i, tid in enumerate(ids):
            task = tasks[tid]
            answer = preds.get(tid) or {}
            outputs = {int(k): (Path(path).parent / v) for k, v in answer.items()} if isinstance(answer, dict) else {}
            rows = score_cases(task, outputs, Path(scratch) / str(i))
            soft = mean(rows, "ok") if rows else 0.0
            items.append({"id": tid, "kind": task.kind, "ok": all(r["ok"] for r in rows) and bool(rows), "soft": soft,
                          "cases": [{k: r[k] for k in ("case", "ok", "detail")} for r in rows],
                          "gold": [{"case": r["case"], "diff": r["gold"]} for r in rows if not r["ok"]]})
            # Reference: return the input spreadsheet unchanged (some answer regions already hold partially correct values)
            reference.append(mean([{"ok": compare(g, p, task.meta["answer_position"])[0]}
                                   for g, p in zip(task.answer, task.meta["inputs"])], "ok"))
    by_kind = {}
    for kind in ("cell", "sheet"):
        rows = [r for r in items if r["kind"] == kind]
        if rows:
            by_kind[kind] = {"n": len(rows), "soft": mean(rows, "soft"), "hard": mean(rows, "ok")}
    return {"metrics": {"soft": mean(items, "soft"), "hard": mean(items, "ok"), "missing": sum(t not in preds for t in ids),
                        "by_kind": by_kind},
            "items": items, "reference": {"unchanged_input": {"soft": sum(reference) / max(1, len(reference))}}}
