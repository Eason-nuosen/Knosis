"""Spreadsheet structure inspection and LibreOffice formula recalculation."""

from __future__ import annotations

import argparse
import json
import tempfile
import time
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from knosis.decision.spreadsheetbench import cells, recalc, regions
from knosis import lang
from knosis.tools import write_result


def show(v):
    return v if isinstance(v, (int, float, str, bool)) or v is None else str(v)


def grid(ws, wv, names) -> list[dict]:
    out = []
    for name in names:
        c, v = ws[name], wv[name]
        if c.value is None and v.value is None:
            continue
        row = {"cell": name, "value": show(v.value)}
        if isinstance(c.value, str) and c.value.startswith("="):
            row["formula"] = c.value
        if c.number_format != "General":
            row["format"] = c.number_format
        out.append(row)
    return out


def corner(ws, rows=30, cols=15) -> list[str]:
    return [f"{get_column_letter(c)}{r}" for r in range(1, min(ws.max_row, rows) + 1)
            for c in range(1, min(ws.max_column, cols) + 1)]


def inspect(path, rng=None) -> dict:
    wb = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        formulas = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
        sheets.append({"sheet": ws.title, "dimensions": ws.dimensions, "max_row": ws.max_row, "max_column": ws.max_column,
                       "merged": [str(m) for m in ws.merged_cells.ranges][:50], "formulas": formulas,
                       "data_validations": len(ws.data_validations.dataValidation),
                       "conditional_formats": len(ws.conditional_formatting), "state": ws.sheet_state})
    if rng:
        view = {f"{s}!{r}": grid(wb[s], wv[s], cells(r, wb[s].max_row)) for s, r in regions(rng, wb.sheetnames[0]) if s in wb.sheetnames}
    else:
        view = {ws.title: grid(ws, wv[ws.title], corner(ws)) for ws in wb.worksheets}
    return {"file": str(path), "sheets": sheets, "cells": view}


def main(argv=None):
    p = argparse.ArgumentParser(prog="sheet")
    p.add_argument("command", choices=["inspect", "recalc"])
    p.add_argument("file")
    p.add_argument("--range")
    p.add_argument("--out", required=True)
    args = p.parse_args(argv)
    started, out = time.monotonic(), Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    units = []
    if args.command == "inspect":
        result = inspect(args.file, args.range)
        for s in result["sheets"]:
            units.append({"id": f"sheet/{Path(args.file).stem}/{s['sheet']}", "form": "schema", "status": "observed",
                          "content": lang.t(f"{s['sheet']}: used range {s['dimensions']}, {len(s['merged'])} merged cells, {s['formulas']} formulas",
                                            f"{s['sheet']}：已用范围 {s['dimensions']}，合并单元格 {len(s['merged'])} 处，公式 {s['formulas']} 个"),
                          "provenance": f"tool:sheet@0.1(inspect,file={args.file})", "meta": s})
    else:
        with tempfile.TemporaryDirectory() as tmp:
            fixed = recalc([Path(args.file)], tmp).get(Path(args.file))
            if fixed is None:
                raise SystemExit("LibreOffice recalculation failed")
            result = inspect(fixed, args.range) | {"file": args.file, "recalculated": True}
    (out / f"{args.command}.json").write_text(json.dumps(result, ensure_ascii=False, indent=1, default=str) + "\n",
                                              encoding="utf-8")
    write_result(out, "sheet", vars(args), units, started)
    print(json.dumps(result, ensure_ascii=False, default=str)[:6000])


if __name__ == "__main__":
    main()
